import tkinter as tk
import webbrowser
from tkinter import messagebox, filedialog
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from config_handler import read_default_passwd, write_default_passwd, generate_name_file, get_next_number, read_excel_file, write_excel_file, excel_reader, write_excel
# Global variable to store the selected file path
selected_file_path = None

def create_file(username_entry, passwd_entry, email_entry, default_passwd_var, alpha_var, omega_var):
    base = "alpha" if alpha_var.get() else "omega"
    number = get_next_number(base)
    server = base
    filename = generate_name_file(base, number)
    username = username_entry.get()
    if not username:
        messagebox.showwarning("Error", "Username is required.")
        return

    passwd = passwd_entry.get() if not default_passwd_var.get() else read_default_passwd()
    server = 7 if base == "alpha" else 6

    with open(filename, 'w') as file:
        file.write("[AutoLogin]\nAutoLogin = true\n")
        file.write(f"Password = '{passwd}'\nPassword2 = ''\n")
        file.write(f"Server = {server}\nTamerSlot = 1\n")
        file.write(f"Username = '{username}'\n")
    
    messagebox.showinfo("Info", f"The file {filename} has been created.")

def passwd_toggle_entry(passwd_entry, default_passwd_var):
    passwd_entry.config(state=tk.DISABLED if default_passwd_var.get() else tk.NORMAL)

def toggle_alpha_omega(var, other_var):
    if var.get():
        other_var.set(0)

def set_new_passwd(new_default_passwd_entry):
    new_passwd = new_default_passwd_entry.get()
    if not new_passwd:
        messagebox.showwarning("Error", "New default password required.")
        return
    write_default_passwd(new_passwd)
    messagebox.showinfo("Success", "Default password has been updated.")

def entry(frame, row, column, columspan, padx=10, pady=10, fg="black", bg="#e0f7fa"):
    entry = tk.Entry(frame, font=("Helvetica", 12), fg=fg, bg=bg)
    entry.grid(row=row, column=column, columnspan=columspan, padx=padx, pady=pady)
    return entry

def label(frame, text, row, column, padx=20, pady=10, sticky="w", bg="white", fg="blue"):
    label = tk.Label(frame, text=text, font=("Helvetica", 12), bg=bg, fg=fg)
    label.grid(row=row, column=column, padx=padx, pady=pady, sticky=sticky)
    return label

def check_button(frame, text, variable, on, off, command, row, column, columnspan, padx=20, pady=10, sticky="w", bg="white", fg="blue" ):
    check_bt = tk.Checkbutton(frame, text=text, variable=variable, onvalue=on, offvalue=off, font=("Helvetica", 12), bg=bg, fg=fg, command=command)
    check_bt.grid(row=row, column=column, columnspan=columnspan, padx=padx, pady=pady, sticky=sticky)
    return check_bt

def change_pwd(root):
    change_pwds = tk.Toplevel(root)
    change_pwds.title("Change Default Password")
    change_pwds.geometry("400x200")
    change_pwds.resizable(False, False)
    
    tk.Label(change_pwds, text="Set new default password:", font=("Helvetica", 12)).grid(row=0, column=0, padx=20, pady=10, sticky="w")
    global new_default_passwd_entry
    new_default_passwd_entry = tk.Entry(change_pwds, font=("Helvetica", 12))
    new_default_passwd_entry.grid(row=0, column=1, padx=20, pady=10)
    
    tk.Button(change_pwds, text="Set Default Password", command=lambda: set_new_passwd(new_default_passwd_entry), font=("Helvetica", 12), bg="#4CAF50", fg="white").grid(row=2, column=0, columnspan=2, pady=20)

def create_button(master, text, command, row, column, columnspan=1, padx=5, pady=5, bg='#FFFFFF', fg='black'):
    button = tk.Button(master, text=text, command=command, font=("Helvetica"), bg=bg, fg=fg)
    button.grid(row=row, column=column, columnspan=columnspan, padx=padx, pady=pady, sticky="we")
    return button

def open_website(): 
    webbrowser.open("https://dmo.gameking.com/Sign/SignUp.aspx")
    
def open_github(event=None):
    webbrowser.open("https://github.com/nihanfirdaus62")

def add_entry(username_entry, passwd_entry, email_entry, alpha_var, omega_var, default_passwd_var):
    global selected_file_path
    if not selected_file_path:
        messagebox.showwarning("Error", "No Excel file selected.")
        return
    
    try:    
        df = read_excel_file(selected_file_path)
        if df is None:
            return

        username = username_entry.get().strip()
        passwd = passwd_entry.get().strip()
        email = email_entry.get().strip()
        Server = "alpha" if alpha_var.get() else "omega"
        digi = "enter later"
            
        if not username or not email:
            messagebox.showwarning("Input Error", "Username and Email fields are required.")
            return
            
        if not passwd:
            passwd = read_default_passwd()  
                
        if 'NO' in df.columns: 
            max_no = df['NO'].max() + 1 
        else: max_no = 1
            
        new_entry = pd.DataFrame({
            'NO' : [max_no],
            'USERNAME': [username], 
            'PASSWORD': [passwd], 
            'EMAIL': [email], 
            'SERVER': [Server], 
            'DIGI':[digi], 
            'STATUS': [2]
            })
            
        df = pd.concat([df, new_entry], ignore_index=True)

        df.to_excel(selected_file_path, index=False, sheet_name='Sheet1')

        wb = load_workbook(selected_file_path)
        ws = wb.active


        tab = Table(displayName="Table1", ref=f"A1:H{len(df)+1}")
        style = TableStyleInfo(name="TableStyleMedium28", showFirstColumn=False,
                            showLastColumn=False, showRowStripes=True, showColumnStripes=True)
        tab.tableStyleInfo = style
        ws.add_table(tab)
            
        wb.save(selected_file_path)
        messagebox.showinfo("Success", f"Data written to {selected_file_path}")
            
        clear(username_entry, passwd_entry, email_entry, default_passwd_var)
    except PermissionError:
        messagebox.showerror("Error", "The Excel file is already open. Please close it and try again.") 
        return

def clear(username_entry, passwd_entry, email_entry,default_passwd_var): 
    username_entry.delete(0, tk.END) 
    email_entry.delete(0, tk.END)
    if not default_passwd_var.get(): 
        passwd_entry.delete(0, tk.END)

def select_excel_file():
    global selected_file_path
    selected_file_path = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx")])
    if selected_file_path:
        messagebox.showinfo("File Selected", f"Selected file: {selected_file_path}")
        
def show_author_message(root):
    c_msg = tk.Toplevel(root)
    c_msg.title("Author message")
    c_msg.geometry("300x100")
    c_msg.resizable(False, False)
    
    msg = tk.Label(c_msg, text="Created By Nihan")
    msg.grid(row=1, column=3, columnspan=1,padx=20, sticky="w")
    
    msg_l = tk.Label(c_msg, text="Github Link!",fg="blue", cursor="hand2")
    msg_l.grid(row=2, column=3, columnspan=1,padx=20, sticky="w")
    msg_l.bind("<Button-1>", open_github)
    
